#!/usr/bin/env python3
"""
Monthly website visit report generator.

Run on the 2nd day of each month. By default it generates the previous
month's raw log file, text report, summary JSON, and Word report.

Usage:
  python monthly_visit_report.py
  python monthly_visit_report.py --month 202606
"""

from __future__ import annotations

import argparse
import importlib.util
import io
import json
import time
from collections import Counter, defaultdict
from contextlib import redirect_stdout
from datetime import date, datetime, timedelta
from pathlib import Path

import requests
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
STATS_SCRIPT = BASE_DIR / "stats_visit.py"
FONT = "Microsoft YaHei"


def previous_month(today: date | None = None) -> tuple[int, int]:
    today = today or date.today()
    first_this_month = today.replace(day=1)
    last_prev_month = first_this_month - timedelta(days=1)
    return last_prev_month.year, last_prev_month.month


def month_range(year: int, month: int) -> tuple[str, str, str, str]:
    start = date(year, month, 1)
    if month == 12:
        end = date(year + 1, 1, 1)
    else:
        end = date(year, month + 1, 1)
    label = f"{year}{month:02d}"
    title = f"{year}年{month}月"
    return start.strftime("%Y%m%d"), end.strftime("%Y%m%d"), label, title


def load_stats_module():
    spec = importlib.util.spec_from_file_location("stats_visit", STATS_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def parse_logs(raw_data: str) -> list[tuple[str, str]]:
    records = []
    for line in raw_data.strip().splitlines():
        parts = line.strip().split()
        if len(parts) >= 2:
            records.append((parts[0], parts[-1]))
    return records


def classify_ips_batch(ips: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    chunks = [ips[i : i + 100] for i in range(0, len(ips), 100)]
    for index, chunk in enumerate(chunks, 1):
        try:
            response = requests.post(
                "http://ip-api.com/batch?fields=status,countryCode,query,message",
                json=chunk,
                timeout=30,
            )
            for item in response.json():
                ip = item.get("query")
                country = item.get("countryCode", "") if item.get("status") == "success" else ""
                if country == "CN":
                    category = "Mainland"
                elif country == "HK":
                    category = "HK"
                else:
                    category = "Overseas"
                if ip:
                    result[ip] = category
        except Exception:
            for ip in chunk:
                result[ip] = "Overseas"
        if index < len(chunks):
            time.sleep(2)
    for ip in ips:
        result.setdefault(ip, "Overseas")
    return result


def set_cell_text(cell, text, bold=False, size=9, color="000000", align="left"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table):
    tbl_pr = table._tbl.tblPr
    table_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if table_borders is None:
        table_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(table_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = table_borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            table_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D0D7DE")


def add_table(doc, headers, rows, widths=None, size=8.5, center_cols=None):
    center_cols = center_cols or []
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_text(cell, header, bold=True, size=size, color="0B2545", align="center")
        shade(cell, "E8EEF5")
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            set_cell_text(cells[index], value, size=size, align="center" if index in center_cols else "left")
    doc.add_paragraph()
    return table


def add_paragraph(doc, text, size=10.5):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    return paragraph


def build_word_report(report_path: Path, summary: dict, raw_name: str, title_month: str):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Heading 1", 15, "1F4D78"), ("Heading 2", 12.5, "2E74B5")]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title_run = title.add_run(f"中烟国际集团网站\n{title_month}访问统计报告")
    title_run.font.name = FONT
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title_run.font.size = Pt(22)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(f"统计周期：{summary['period_display']}")
    sub_run.font.name = FONT
    sub_run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor.from_string("555555")

    doc.add_paragraph()
    add_table(
        doc,
        ["项目", "数据"],
        [
            ("访问次数", summary["total_visits"]),
            ("独立 IP 数", summary["unique_ips"]),
            ("原始数据文件", raw_name),
            ("统计生成时间", summary["generated_at"]),
        ],
        widths=[1.7, 4.8],
        size=9.5,
    )

    doc.add_heading("1. 地域分类统计", level=1)
    category_rows = []
    labels = {"Mainland": "中国内地", "HK": "中国香港", "Overseas": "海外/其他"}
    for category in ["Mainland", "HK", "Overseas"]:
        stats = summary["categories"][category]
        percent = stats["visits"] / summary["total_visits"] * 100 if summary["total_visits"] else 0
        category_rows.append((labels[category], stats["visits"], stats["unique_ips"], f"{percent:.1f}%"))
    add_table(doc, ["分类", "访问次数", "独立 IP", "访问占比"], category_rows, widths=[1.7, 1.2, 1.2, 1.2], size=9, center_cols=[1, 2, 3])

    doc.add_heading("2. 访问量 Top 30 IP", level=1)
    top_ip_rows = [
        (item["ip"], labels[item["category"]], item["visits"], item["top_ref"][:45])
        for item in summary["top_ips"]
    ]
    add_table(doc, ["IP", "归属", "访问次数", "主要浏览器/来源字段"], top_ip_rows, widths=[1.55, 1.1, 0.9, 3.05], size=7.8, center_cols=[1, 2])

    doc.add_heading("3. 浏览器/来源字段 Top 30", level=1)
    add_paragraph(doc, "说明：原统计脚本取 nginx 日志行尾字段作为 Referer 输出，实际更接近浏览器或爬虫 User-Agent 片段，以下按脚本口径整理。", size=9.5)
    add_table(doc, ["字段", "次数"], [(key[:70], value) for key, value in summary["top_refs"]], widths=[5.3, 1.0], size=8.2, center_cols=[1])

    doc.add_heading("4. 简要结论", level=1)
    add_paragraph(
        doc,
        f"{title_month}网站统计访问次数为 {summary['total_visits']} 次，独立 IP 为 {summary['unique_ips']} 个。"
        f"按 IP 归属统计，中国内地访问 {summary['categories']['Mainland']['visits']} 次，"
        f"中国香港访问 {summary['categories']['HK']['visits']} 次，"
        f"海外/其他访问 {summary['categories']['Overseas']['visits']} 次。",
    )
    add_paragraph(doc, "日志中可能包含搜索引擎、AI 爬虫或其他自动化访问字段；本报告仅按访问日志整理，不直接代表安全攻击结论。")
    doc.save(report_path)


def generate_report(year: int, month: int):
    start_date, end_date, label, title_month = month_range(year, month)
    raw_path = BASE_DIR / f"raw_logs_{label}.json"
    text_path = BASE_DIR / f"stats_{label}.txt"
    summary_path = BASE_DIR / f"stats_{label}_summary.json"
    report_path = BASE_DIR / f"中烟国际集团网站{year}年{month}月访问统计报告.docx"

    stats = load_stats_module()
    stats.START_DATE = start_date
    stats.END_DATE = end_date
    stats.RAW_DATA_FILE = str(raw_path)

    output = io.StringIO()
    with redirect_stdout(output):
        print("=" * 60)
        print("   Website Visit Stats - CTI Group")
        print("=" * 60)
        print(f"\nDate Range: {start_date} - {end_date}\n")
        raw_data = stats.get_raw_logs()
        if raw_data.strip():
            stats.save_raw_data(raw_data, str(raw_path))
            stats.analyze(raw_data, title_month)
        else:
            print("[ERROR] No data fetched, check connection and date range")
            text_path.write_text(output.getvalue(), encoding="utf-8")
            raise RuntimeError("No data fetched")

    text_path.write_text(output.getvalue(), encoding="utf-8")

    raw_payload = json.loads(raw_path.read_text(encoding="utf-8"))
    records = parse_logs(raw_payload["raw_data"])
    ip_counts = Counter(ip for ip, _ in records)
    ref_counts = Counter(ref for _, ref in records)
    categories_by_ip = classify_ips_batch(list(ip_counts.keys()))

    category_visits = defaultdict(int)
    category_ips = defaultdict(set)
    for ip, count in ip_counts.items():
        category = categories_by_ip[ip]
        category_visits[category] += count
        category_ips[category].add(ip)

    top_ips = []
    for ip, count in ip_counts.most_common(30):
        top_ref = Counter(ref for record_ip, ref in records if record_ip == ip).most_common(1)[0][0]
        top_ips.append({"ip": ip, "category": categories_by_ip[ip], "visits": count, "top_ref": top_ref})

    period_display = f"{year}年{month}月1日 - {year}年{month}月{(date(year + (month // 12), month % 12 + 1, 1) - timedelta(days=1)).day}日"
    summary = {
        "period": label,
        "period_display": period_display,
        "date_range": f"{start_date}-{end_date}",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_visits": len(records),
        "unique_ips": len(ip_counts),
        "categories": {
            category: {"visits": category_visits[category], "unique_ips": len(category_ips[category])}
            for category in ["Mainland", "HK", "Overseas"]
        },
        "top_ips": top_ips,
        "top_refs": ref_counts.most_common(30),
        "files": {
            "raw": raw_path.name,
            "text": text_path.name,
            "summary": summary_path.name,
            "word": report_path.name,
        },
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    build_word_report(report_path, summary, raw_path.name, title_month)
    return summary


def parse_args():
    parser = argparse.ArgumentParser(description="Generate monthly website visit report.")
    parser.add_argument("--month", help="Month to generate, format YYYYMM. Default: previous month.")
    return parser.parse_args()


def main():
    args = parse_args()
    if args.month:
        if len(args.month) != 6 or not args.month.isdigit():
            raise SystemExit("--month must be YYYYMM, for example 202606")
        year = int(args.month[:4])
        month = int(args.month[4:])
    else:
        year, month = previous_month()

    summary = generate_report(year, month)
    print(json.dumps(summary["files"], ensure_ascii=False, indent=2))
    print(f"Done: {summary['period_display']}")


if __name__ == "__main__":
    main()
